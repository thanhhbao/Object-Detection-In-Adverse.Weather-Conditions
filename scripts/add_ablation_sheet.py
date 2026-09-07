#!/usr/bin/env python3
"""Add an ablation-results tab to the metrics workbook.

Four blocks on one sheet:
  1. Official test metrics for Phase 2 / A0R / A1-DINO v1, read straight from
     the official_eval backup JSON so nothing is transcribed by hand.
  2. Per-class F1 on val for Phase 2 vs A1-DINO v2.
  3. Per-class confidence thresholds fitted on val.
  4. Whether those thresholds carried over to test — the number to report.

Usage:
  python scripts/add_ablation_sheet.py \
    --workbook docs/adverse_weather_model_results_google_sheets.xlsx \
    --metrics-root <backup>/metrics \
    --results <staged>/ablation_results.json
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET = "Ablation — Retrieval + Conf"
CLASSES = ["person", "bicycle", "car", "motorcycle", "bus", "truck"]
RARE = {"bicycle", "motorcycle", "bus"}
LOW_SUPPORT = 100
DATASETS = [("xwod", "XWOD"), ("bdd", "BDD"), ("dawn", "DAWN"), ("acdc", "ACDC")]

TITLE = Font(bold=True, size=13)
HEAD = Font(bold=True, size=10, color="FFFFFF")
BOLD = Font(bold=True, size=10)
BASE = Font(size=10)
DIM = Font(size=9, italic=True, color="666666")
HEAD_FILL = PatternFill("solid", fgColor="44546A")
BAND = PatternFill("solid", fgColor="F2F2F2")
GOOD = PatternFill("solid", fgColor="E2EFDA")
BAD = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def f1_of(m: dict) -> float:
    p, r = m["precision"], m["recall"]
    return 2 * p * r / (p + r) if (p + r) else 0.0


def put(ws, row, col, value, font=BASE, fill=None, fmt=None, align="left"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.border = THIN
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def header(ws, row, labels, start=1):
    for i, h in enumerate(labels):
        put(ws, row, start + i, h, HEAD, HEAD_FILL,
            align="center" if i else "left")
    return row + 1


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", type=Path, required=True)
    ap.add_argument("--metrics-root", type=Path, required=True)
    ap.add_argument("--results", type=Path, required=True)
    args = ap.parse_args()

    D = json.loads(args.results.read_text(encoding="utf-8"))

    official = {}
    for arm in ("phase2", "a0r", "a1_dino"):
        for ds, _ in DATASETS:
            p = args.metrics_root / arm / f"official_{ds}_test_metrics.json"
            if p.is_file():
                official[(arm, ds)] = json.loads(p.read_text(encoding="utf-8"))

    missing = [f"{a}/{d}" for a in ("phase2", "a0r", "a1_dino")
               for d, _ in DATASETS if (a, d) not in official]
    if missing:
        raise SystemExit(
            f"Missing official metrics under {args.metrics_root}: "
            + ", ".join(missing)
            + "\nExtract official_eval_*.tar.gz first; a silently empty block "
              "would look like a result.")

    wb = openpyxl.load_workbook(args.workbook)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)

    r = 1
    put(ws, r, 1, "Ablation — truy hồi dữ liệu và ngưỡng theo lớp (RT-DETR-L)", TITLE)
    r += 1
    put(ws, r, 1, "Ba cấu hình dùng chung kiến trúc và cùng 55.985 ảnh huấn luyện; "
                  "chỉ khác cách chọn 5.000 ảnh bổ sung.", DIM)
    r += 2

    # ── Block 1: official test metrics, v1 generation ───────────────────────
    put(ws, r, 1, "1 — Kết quả trên tập test (Phase 2 / A0R / A1-DINO v1)", BOLD)
    r += 1
    put(ws, r, 1, "Nguồn: official_eval_20260901T180820Z. F1 = 2PR/(P+R).", DIM)
    r += 1
    r = header(ws, r, ["Tập test", "Cấu hình", "Precision", "Recall", "F1",
                       "mAP@50", "mAP@50:95", "Δ F1 vs Phase 2"])
    for ds, dname in DATASETS:
        if ("phase2", ds) not in official:
            continue
        base = f1_of(official[("phase2", ds)])
        for arm, aname in (("phase2", "Phase 2"), ("a0r", "A0R"),
                           ("a1_dino", "A1-DINO v1")):
            m = official.get((arm, ds))
            if not m:
                continue
            f = f1_of(m)
            band = BAND if ds in ("bdd", "acdc") else None
            put(ws, r, 1, dname, BASE, band)
            put(ws, r, 2, aname, BASE, band)
            for i, v in enumerate([m["precision"], m["recall"], f,
                                   m["map50"], m["map50_95"]]):
                put(ws, r, 3 + i, round(v, 4), BASE, band, "0.0000", "center")
            d = f - base
            put(ws, r, 8, "—" if arm == "phase2" else round(d, 4),
                BASE, band if arm == "phase2" else (GOOD if d > 0 else BAD),
                None if arm == "phase2" else "+0.0000;-0.0000", "center")
            r += 1
    r += 1

    # ── Block 2: per-class F1 on val, v2 ────────────────────────────────────
    put(ws, r, 1, "2 — F1 theo lớp trên tập val (Phase 2 vs A1-DINO v2)", BOLD)
    r += 1
    put(ws, r, 1, "A1-DINO v2 sửa ba lỗi của v1: nhúng theo vùng cắt thay vì "
                  "toàn ảnh, quota theo lớp, ACDC nhân đôi trọng số truy vấn. "
                  "⚠ = dưới 100 hộp, F1 quá nhiễu để xếp hạng.", DIM)
    r += 1
    r = header(ws, r, ["Tập val", "Lớp", "Số hộp", "Phase 2", "A1-DINO v2",
                       "Δ", "Ghi chú"])
    for ds, dname in DATASETS:
        blk = D["val_f1"].get(ds)
        if not blk:
            continue
        for c in CLASSES + ["ALL"]:
            sup = blk["support"][c]
            a, b = blk["phase2"][c], blk["a1_dino_v2"][c]
            d = b - a
            is_all = c == "ALL"
            fnt = BOLD if is_all else BASE
            fill = BAND if is_all else None
            put(ws, r, 1, dname if is_all else "", fnt, fill)
            put(ws, r, 2, "TỔNG" if is_all else
                (f"{c} ⚠" if 0 < sup < LOW_SUPPORT else c), fnt, fill)
            put(ws, r, 3, sup, fnt, fill, "#,##0", "center")
            put(ws, r, 4, round(a, 4), fnt, fill, "0.0000", "center")
            put(ws, r, 5, round(b, 4), fnt, fill, "0.0000", "center")
            put(ws, r, 6, round(d, 4), fnt,
                GOOD if d > 0 else (BAD if d < 0 else fill),
                "+0.0000;-0.0000", "center")
            note = ""
            if is_all:
                note = "A1-DINO v2 thắng" if d > 0 else "Phase 2 thắng"
            elif 0 < sup < LOW_SUPPORT:
                note = "support thấp — không kết luận"
            elif c in RARE:
                note = "lớp hiếm"
            put(ws, r, 7, note, DIM if note else BASE, fill)
            r += 1
    r += 1

    # ── Block 3: thresholds fitted on val ───────────────────────────────────
    put(ws, r, 1, "3 — Ngưỡng tin cậy theo lớp, dò trên val (A1-DINO v2)", BOLD)
    r += 1
    put(ws, r, 1, "Ultralytics đọc mọi lớp tại một ngưỡng chung tối đa hoá F1 "
                  "trung bình. Cột 'riêng' là ngưỡng tối ưu của từng lớp.", DIM)
    r += 1
    r = header(ws, r, ["Tập val", "Lớp", "Số hộp", "Ngưỡng chung", "Ngưỡng riêng",
                       "F1 chung", "F1 riêng", "Δ"])
    for ds, dname in DATASETS:
        blk = D["thr_val"].get(ds)
        if not blk:
            continue
        for c in CLASSES:
            sup, t, fs, ft = blk["cls"][c]
            d = ft - fs
            put(ws, r, 1, "")
            put(ws, r, 2, f"{c} ⚠" if 0 < sup < LOW_SUPPORT else c)
            put(ws, r, 3, sup, BASE, None, "#,##0", "center")
            put(ws, r, 4, blk["shared_t"], BASE, None, "0.00", "center")
            put(ws, r, 5, t, BASE, None, "0.00", "center")
            put(ws, r, 6, round(fs, 4), BASE, None, "0.0000", "center")
            put(ws, r, 7, round(ft, 4), BASE, None, "0.0000", "center")
            put(ws, r, 8, round(d, 4), BASE,
                GOOD if d > 0 else None, "+0.0000;-0.0000", "center")
            r += 1
        put(ws, r, 1, dname, BOLD, BAND)
        put(ws, r, 2, "F1 trung bình", BOLD, BAND)
        put(ws, r, 3, "", BOLD, BAND)
        put(ws, r, 4, blk["shared_t"], BOLD, BAND, "0.00", "center")
        put(ws, r, 5, "", BOLD, BAND)
        put(ws, r, 6, round(blk["mean_shared"], 4), BOLD, BAND, "0.0000", "center")
        put(ws, r, 7, round(blk["mean_tuned"], 4), BOLD, BAND, "0.0000", "center")
        put(ws, r, 8, round(blk["gain"], 4), BOLD, GOOD, "+0.0000;-0.0000", "center")
        r += 1
    r += 1

    # ── Block 4: carryover to test — the number that counts ─────────────────
    put(ws, r, 1, "4 — Ngưỡng dò trên val, áp lên test: mức tăng có sống sót không", BOLD)
    r += 1
    put(ws, r, 1, "Đây mới là con số được phép báo cáo. Ngưỡng cố định từ val "
                  "trước khi nhìn test, nên đây là một lần dùng test hợp lệ.", DIM)
    r += 1
    r = header(ws, r, ["Tập test", "Lớp", "Số hộp", "Ngưỡng", "F1 chung",
                       "F1 riêng", "Δ trên test", "Δ khi dò (val)", "Giữ lại"])
    for ds, dname in DATASETS:
        blk = D["thr_test"].get(ds)
        if not blk:
            put(ws, r, 1, dname, BASE, WARN)
            put(ws, r, 2, "chưa chạy", DIM, WARN)
            for cc in range(3, 10):
                put(ws, r, cc, "", BASE, WARN)
            r += 1
            continue
        for c in CLASSES:
            sup, t, fs, ft = blk["cls"][c]
            d = ft - fs
            put(ws, r, 1, "")
            put(ws, r, 2, f"{c} ⚠" if 0 < sup < LOW_SUPPORT else c)
            put(ws, r, 3, sup, BASE, None, "#,##0", "center")
            put(ws, r, 4, t, BASE, None, "0.00", "center")
            put(ws, r, 5, round(fs, 4), BASE, None, "0.0000", "center")
            put(ws, r, 6, round(ft, 4), BASE, None, "0.0000", "center")
            put(ws, r, 7, round(d, 4), BASE,
                GOOD if d > 0 else (BAD if d < 0 else None),
                "+0.0000;-0.0000", "center")
            put(ws, r, 8, "")
            put(ws, r, 9, "")
            r += 1
        fitted = D["thr_val"][ds]["gain"]
        gain = blk["gain"]
        ratio = gain / fitted if fitted else None
        put(ws, r, 1, dname, BOLD, BAND)
        put(ws, r, 2, "F1 trung bình", BOLD, BAND)
        for cc in (3, 4):
            put(ws, r, cc, "", BOLD, BAND)
        put(ws, r, 5, round(blk["mean_shared"], 4), BOLD, BAND, "0.0000", "center")
        put(ws, r, 6, round(blk["mean_tuned"], 4), BOLD, BAND, "0.0000", "center")
        put(ws, r, 7, round(gain, 4), BOLD, GOOD if gain > 0 else BAD,
            "+0.0000;-0.0000", "center")
        put(ws, r, 8, round(fitted, 4), BOLD, BAND, "+0.0000;-0.0000", "center")
        put(ws, r, 9, "âm" if gain <= 0 else f"{ratio:.0%}", BOLD,
            BAD if gain <= 0 else (GOOD if ratio >= 0.5 else WARN), None, "center")
        r += 1
    r += 1

    # ── Notes ───────────────────────────────────────────────────────────────
    put(ws, r, 1, "Ghi chú", BOLD)
    r += 1
    for note in [
        "Mỗi cấu hình chỉ chạy một seed, không có khoảng tin cậy. Biên độ giữa "
        "các cấu hình nhỏ hơn nhiều so với sai số mong đợi của một lần chạy.",
        "A1-DINO v1 được đánh giá trên test, kết quả đó dẫn tới việc sửa phương "
        "pháp và huấn luyện v2. Vì vậy số test của v2 đã qua một vòng lặp có "
        "thông tin test, không còn là ước lượng giữ lại thuần tuý.",
        "Ngưỡng theo lớp chỉ có lợi khi phân phối điểm giữa các lớp lệch nhau rõ "
        "rệt (ACDC giữ 48%). Khi đường cong F1 đã phẳng quanh đỉnh thì mức tăng "
        "đo trên val là nhiễu và không chuyển giao được (XWOD âm).",
        "Nguồn bổ sung là BDD100K — ảnh đô thị trời quang. Toàn bộ pool chỉ có "
        "1.391 ảnh chứa motorcycle và đã lấy hết, nên với lớp này truy hồi suy "
        "biến thành lấy toàn bộ.",
    ]:
        put(ws, r, 1, "• " + note, DIM)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        r += 1

    for col, w in zip("ABCDEFGHI", [16, 22, 11, 14, 14, 12, 12, 16, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    wb.save(args.workbook)
    print(f"Sheet '{SHEET}' → {args.workbook}")
    print(f"Sheets: {wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
